import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


SCRIPTS_DIR = Path(__file__).resolve().parents[1] / 'scripts'
sys.path.insert(0, str(SCRIPTS_DIR))

from export_source_sheets import export_source_sheets  # noqa: E402


class SourceSheetExportTests(unittest.TestCase):
    def test_existing_sheets_are_preserved_and_source_sheets_are_added(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            source = temp / 'existing.xlsx'
            output = temp / 'with-source.xlsx'
            history = temp / 'history.json'

            workbook = Workbook()
            workbook.active.title = 'Volume'
            workbook['Volume']['A1'] = 'existing-value'
            workbook.save(source)

            history.write_text(json.dumps([{
                'date': '2026-09-26',
                'brand_sov_by_source': {
                    'ppomppu': {'유모바일': 1},
                    'dc': {'유모바일': 2},
                },
                'keywords_raw_by_source': {
                    'ppomppu': {'추천인': 1},
                    'dc': {'추천인': 2},
                },
            }], ensure_ascii=False), encoding='utf-8')

            export_source_sheets(source, history, output)
            result = load_workbook(output, data_only=True)

            self.assertEqual('existing-value', result['Volume']['A1'].value)
            self.assertEqual(
                ['date', 'source', 'brand', 'mentions'],
                [cell.value for cell in result['Brand_SOV_Source_Long'][1]],
            )
            self.assertEqual(
                ('2026-09-26', 'ppomppu', '유모바일', 1),
                tuple(cell.value for cell in result['Brand_SOV_Source_Long'][2]),
            )
            self.assertIn('Keywords_Source_Raw', result.sheetnames)
            self.assertIn('Keywords_Source_Normalized', result.sheetnames)

    def test_input_workbook_cannot_be_overwritten_directly(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            workbook_path = temp / 'existing.xlsx'
            history_path = temp / 'history.json'
            Workbook().save(workbook_path)
            history_path.write_text('[]', encoding='utf-8')

            with self.assertRaises(ValueError):
                export_source_sheets(workbook_path, history_path, workbook_path)


if __name__ == '__main__':
    unittest.main()
