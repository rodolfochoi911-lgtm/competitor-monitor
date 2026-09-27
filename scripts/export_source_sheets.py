"""Add source-level sheets to a copy of an existing MVNO workbook."""

import argparse
import json
from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook
from openpyxl.styles import Font


SOURCE_SHEETS = {
    'Brand_SOV_Source_Long': ['date', 'source', 'brand', 'mentions'],
    'Keywords_Source_Raw': ['date', 'source', 'keyword', 'count'],
    'Keywords_Source_Normalized': ['date', 'source', 'term', 'type', 'count'],
}


def load_history(path):
    with Path(path).open(encoding='utf-8') as handle:
        history = json.load(handle)
    if not isinstance(history, list):
        raise ValueError('dashboard history must be a JSON list')
    return history


def source_rows(history):
    brand_rows = []
    keyword_rows = []
    normalized_rows = []

    aliases = {
        alias.lower()
        for keywords in __import__('monitor_crawler').BRAND_KEYWORDS.values()
        for alias in keywords
    }

    for entry in history:
        date = entry.get('date', '')
        brands_by_source = entry.get('brand_sov_by_source', {})
        keywords_by_source = entry.get('keywords_raw_by_source', {})

        for source in ('ppomppu', 'dc'):
            brand_counts = brands_by_source.get(source, {})
            keyword_counts = keywords_by_source.get(source, {})

            for brand, count in sorted(brand_counts.items()):
                count = int(count)
                brand_rows.append([date, source, brand, count])
                normalized_rows.append([date, source, brand, 'brand', count])

            for keyword, count in sorted(keyword_counts.items()):
                count = int(count)
                keyword_rows.append([date, source, keyword, count])
                if keyword.lower() not in aliases:
                    normalized_rows.append([date, source, keyword, 'keyword', count])

    return brand_rows, keyword_rows, normalized_rows


def replace_sheet(workbook, title, headers, rows):
    if title in workbook.sheetnames:
        workbook.remove(workbook[title])

    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions


def export_source_sheets(workbook_path, history_path, output_path):
    workbook_path = Path(workbook_path)
    output_path = Path(output_path)
    if workbook_path.resolve() == output_path.resolve():
        raise ValueError('output must be a different file; validate it before replacing production')

    output_path.parent.mkdir(parents=True, exist_ok=True)
    copy2(workbook_path, output_path)

    keep_vba = output_path.suffix.lower() == '.xlsm'
    workbook = load_workbook(output_path, keep_vba=keep_vba, keep_links=True)
    brand_rows, keyword_rows, normalized_rows = source_rows(load_history(history_path))

    rows_by_sheet = {
        'Brand_SOV_Source_Long': brand_rows,
        'Keywords_Source_Raw': keyword_rows,
        'Keywords_Source_Normalized': normalized_rows,
    }
    for title, headers in SOURCE_SHEETS.items():
        replace_sheet(workbook, title, headers, rows_by_sheet[title])

    workbook.save(output_path)
    return rows_by_sheet


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--workbook', required=True, help='Existing .xlsx or .xlsm workbook')
    parser.add_argument('--history', default='data/dashboard_history.json')
    parser.add_argument('--output', required=True, help='New workbook path; must differ from input')
    args = parser.parse_args()

    rows = export_source_sheets(args.workbook, args.history, args.output)
    counts = ', '.join(f'{name}={len(values)}' for name, values in rows.items())
    print(f'Created {args.output}: {counts}')


if __name__ == '__main__':
    main()
